"""Preview and atomically import user-provided Excel/CSV order snapshots.

Uploads stay in memory. Missing source tables mean unknown coverage, not zero receipts.
An import replaces the selected imported orders' snapshot; unrelated orders/runs survive.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import threading
import time
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from supplylens.db import (
    SHANGHAI,
    ETAHistory,
    LogisticsEvent,
    OrderOrigin,
    PurchaseOrder,
    Receipt,
    Shipment,
    Supplier,
    cutoff_after,
    local_datetime,
    validate_seed,
)

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TOTAL_BYTES = 20 * 1024 * 1024
MAX_ROWS = 10000
MAX_PREVIEWS = 20
PREVIEW_TTL = 15 * 60
TABLE_FIELDS = {
    "suppliers": ("id", "name"),
    "orders": ("id", "supplier_id", "material_code", "qty", "promised_date", "created_at"),
    "shipments": ("id", "order_id", "qty", "shipped_at"),
    "eta_history": ("id", "shipment_id", "eta", "updated_at"),
    "logistics_events": ("id", "shipment_id", "status", "event_time", "note"),
    "receipts": ("id", "shipment_id", "qty", "received_at"),
}
TABLE_NAMES = {
    "suppliers": "suppliers", "供应商": "suppliers",
    "orders": "orders", "采购订单": "orders", "采购单": "orders",
    "shipments": "shipments", "发运单": "shipments", "发运": "shipments",
    "eta_history": "eta_history", "预计到货": "eta_history", "到货预测": "eta_history",
    "logistics_events": "logistics_events", "物流事件": "logistics_events", "物流": "logistics_events",
    "receipts": "receipts", "收货记录": "receipts", "收货": "receipts",
}
ALIASES = {
    "supplier_id": ("供应商编号", "供应商编码"), "order_id": ("采购单号", "订单编号", "订单号"),
    "shipment_id": ("发运单号", "发运编号", "批次编号"), "material_code": ("物料编码", "物料编号"),
    "qty": ("数量", "采购数量", "订购数量", "发运数量", "收货数量", "已收数量"),
    "promised_date": ("承诺到货日", "承诺日期"), "created_at": ("创建时间", "下单时间"),
    "shipped_at": ("发运时间", "发货时间"), "eta": ("预计到货日", "预计到货日期", "eta日期"),
    "updated_at": ("更新时间", "预测更新时间"), "status": ("物流状态", "状态"),
    "event_time": ("事件时间", "物流时间"), "note": ("备注", "物流备注"),
    "received_at": ("收货时间", "入库时间"), "name": ("供应商名称", "名称"),
}
ID_ALIASES = {
    "suppliers": ("供应商编号", "供应商编码"), "orders": ("采购单号", "订单编号", "订单号"),
    "shipments": ("发运单号", "发运编号", "批次编号"),
    "eta_history": ("预测编号", "eta编号"), "logistics_events": ("事件编号", "物流事件编号"),
    "receipts": ("收货编号", "收货记录编号"),
}
STATUSES = {"已发运": "departed", "在途": "in_transit", "运输中": "in_transit", "延误": "delayed",
            "滞留": "held", "清关滞留": "customs_hold", "异常": "exception", "已签收": "delivered", "已到达": "arrived"}
MODELS = {"suppliers": Supplier, "orders": PurchaseOrder, "shipments": Shipment, "eta_history": ETAHistory,
          "logistics_events": LogisticsEvent, "receipts": Receipt}
DATE_FIELDS = {"promised_date", "eta"}
DATETIME_FIELDS = {"created_at", "shipped_at", "updated_at", "event_time", "received_at"}


class ImportValidationError(ValueError):
    def __init__(self, errors: list[str]):
        self.errors = errors[:30]
        super().__init__("；".join(self.errors))


class ImportConflict(ImportValidationError):
    pass


def _header(value) -> str:
    return str(value or "").strip().lower()


def _field_map(table: str) -> dict[str, str]:
    result = {field: field for field in TABLE_FIELDS[table]}
    for field in TABLE_FIELDS[table]:
        names = ID_ALIASES[table] if field == "id" else ALIASES.get(field, ())
        result.update({_header(name): field for name in names})
    return result


def _value(field: str, value):
    if isinstance(value, bool):
        raise ValueError("不能使用布尔值")
    if value is None or value == "":
        if field == "note":
            return ""
        raise ValueError("必填值为空")
    if isinstance(value, str) and value.lstrip().startswith("="):
        raise ValueError("不接受公式，请粘贴为值")
    if field == "qty":
        raw_number = str(value)
        if "," in raw_number and not re.fullmatch(r"[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?", raw_number):
            raise ValueError("千位分隔符格式不正确")
        try:
            number = Decimal(raw_number.replace(",", ""))
        except InvalidOperation:
            raise ValueError("数量必须为正整数") from None
        if not number.is_finite() or number != number.to_integral_value() or not 0 < number <= 2147483647:
            raise ValueError("数量必须为正整数且不超过数据库范围")
        return int(number)
    if field in DATE_FIELDS:
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return date.fromisoformat(str(value).strip().replace("/", "-")).isoformat()
    if field in DATETIME_FIELDS:
        if isinstance(value, date) and not isinstance(value, datetime):
            value = datetime.combine(value, datetime.min.time())
        normalized = local_datetime(value if isinstance(value, datetime) else str(value).strip().replace("/", "-"))
        return normalized.replace(tzinfo=SHANGHAI).isoformat()
    if field in {"id", "supplier_id", "order_id", "shipment_id", "material_code"} and not isinstance(value, str):
        raise ValueError("编号必须存为文本，以保留前导零和完整精度")
    text = str(value).strip()
    if not text and field != "note":
        raise ValueError("必填值为空")
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", text):
        raise ValueError("包含不支持的控制字符")
    return STATUSES.get(text, text.lower()) if field == "status" else text


def _rows(table: str, rows: list[list], errors: list[str]) -> list[dict]:
    if not rows:
        errors.append(f"{table}：缺少表头")
        return []
    aliases = _field_map(table)
    columns = [aliases.get(_header(value)) for value in rows[0]]
    nonempty = [key for key in columns if key]
    if len(nonempty) != len(set(nonempty)):
        errors.append(f"{table}：存在重复列或映射到同一字段的列")
    missing = set(TABLE_FIELDS[table]) - set(nonempty) - {"note"}
    if missing:
        errors.append(f"{table}：缺少必填列 {', '.join(sorted(missing))}")
    unknown = [str(value) for value, key in zip(rows[0], columns, strict=True) if value not in (None, "") and not key]
    if unknown:
        errors.append(f"{table}：无法识别列 {', '.join(unknown)}；请使用模板字段")
    if missing or unknown or len(nonempty) != len(set(nonempty)):
        return []
    parsed = []
    for row_number, row in enumerate(rows[1:], 2):
        if all(value is None or value == "" for value in row):
            continue
        record = {}
        for position, key in enumerate(columns):
            if key:
                try:
                    record[key] = _value(key, row[position] if position < len(row) else None)
                except (ValueError, TypeError, OverflowError):
                    errors.append(f"{table} 第 {row_number} 行 / {key}：值为空、格式不正确或不是文本编号")
        record.setdefault("note", "") if table == "logistics_events" else None
        parsed.append(record)
        if len(errors) >= 30:
            break
    return parsed


def parse_uploads(files: list[tuple[str, bytes]]) -> tuple[dict, dict, list[str]]:
    errors, warnings, raw = [], [], {}
    if not files or len(files) > 6 or sum(len(body) for _, body in files) > MAX_TOTAL_BYTES:
        raise ImportValidationError(["请选择一个 Excel 或最多六个 CSV，合计不超过 20 MB。"])
    for filename, body in files:
        suffix = Path(filename).suffix.lower()
        if not body or len(body) > MAX_FILE_BYTES:
            errors.append(f"{Path(filename).name}：文件为空或超过 5 MB")
            continue
        if suffix == ".xlsx":
            if len(files) != 1:
                raise ImportValidationError(["Excel 工作簿请单独上传；CSV 可以多选。"])
            try:
                with ZipFile(io.BytesIO(body)) as archive:
                    members = archive.infolist()
                    if len(members) > 2000 or sum(info.file_size for info in members) > 40 * 1024 * 1024:
                        raise ValueError("工作簿解压后过大")
                    for info in members:
                        if info.filename.endswith(".xml"):
                            text = archive.read(info).upper()
                            if b"<!DOCTYPE" in text or b"<!ENTITY" in text:
                                raise ValueError("不支持自定义 XML 实体")
                workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False, keep_links=False)
                try:
                    for sheet in workbook:
                        if sheet.title in {"说明", "Instructions", "README"}:
                            continue
                        table = TABLE_NAMES.get(_header(sheet.title))
                        if not table or table in raw:
                            errors.append(f"无法识别或重复工作表：{sheet.title}")
                            continue
                        data = []
                        for row in sheet.iter_rows():
                            if len(data) > MAX_ROWS or len(row) > 30:
                                raise ValueError("行列数超过上限")
                            if any(cell.data_type == "f" for cell in row):
                                raise ValueError("请将所有公式转换为值后上传")
                            data.append([cell.value for cell in row])
                        raw[table] = data
                finally:
                    workbook.close()
            except (BadZipFile, KeyError, ValueError, OSError, TypeError, OverflowError, ParseError, InvalidFileException):
                errors.append("Excel 文件损坏、含公式、XML 实体或超出允许大小；只支持 .xlsx 数据表。")
        elif suffix == ".csv":
            table = TABLE_NAMES.get(_header(Path(filename).stem))
            if not table or table in raw:
                errors.append(f"CSV 文件名无法识别或重复：{Path(filename).name}")
                continue
            try:
                try:
                    text = body.decode("utf-8-sig")
                except UnicodeDecodeError:
                    text = body.decode("gb18030")
                reader = csv.reader(io.StringIO(text), strict=True)
                data = []
                for row in reader:
                    if len(data) > MAX_ROWS or len(row) > 30:
                        raise ValueError("CSV 行列数超过上限")
                    data.append(row)
                raw[table] = data
            except (UnicodeError, csv.Error, ValueError):
                errors.append(f"{Path(filename).name}：无法读取 CSV，请使用 UTF-8/GB18030 编码和逗号分隔。")
        else:
            errors.append("只支持 .xlsx 和 .csv 文件。")
    for table in ("suppliers", "orders"):
        if table not in raw:
            errors.append(f"必须提供 {table} 表；可下载 Excel 或 CSV 模板。")
    payload = {table: _rows(table, rows, errors) if (rows := raw.get(table)) is not None else []
               for table in TABLE_FIELDS}
    if not errors:
        try:
            validate_seed(payload)
            for table, rows in payload.items():
                if len({row["id"].casefold() for row in rows}) != len(rows):
                    raise ValueError(f"{table} 编号重复或仅大小写不同")
        except ValueError as exc:
            errors.append(f"数据关联或数量校验失败：{exc}")
    if not payload["orders"]:
        errors.append("orders 表至少需要一张采购单。")
    if len(payload["orders"]) > 500:
        errors.append("单批最多导入 500 张订单，请拆分文件。")
    if errors:
        raise ImportValidationError(errors)
    coverage = {table: table in raw for table in TABLE_FIELDS}
    for table in ("shipments", "eta_history", "logistics_events", "receipts"):
        if not coverage[table]:
            warnings.append(f"未提供 {table}，相关信息将标记为未知。若确认没有记录，请保留只有表头的空表。")
    return payload, coverage, warnings


@dataclass
class PreparedImport:
    payload: dict
    coverage: dict
    as_of: date
    source_name: str
    expires_at: float
    fingerprint: str
    warnings: list[str]


class ImportService:
    def __init__(self, engine):
        self.engine = engine
        self._previews: dict[str, PreparedImport] = {}
        self._lock = threading.RLock()

    def _check_database(self, data: dict) -> dict:
        order_ids = {row["id"] for row in data["orders"]}
        with Session(self.engine) as session:
            for table, model in MODELS.items():
                existing_ids = {identifier.casefold(): identifier for identifier in session.scalars(select(model.id)).all()}
                for row in data[table]:
                    if row["id"].casefold() in existing_ids and existing_ids[row["id"].casefold()] != row["id"]:
                        raise ImportConflict([f"{table} 编号与现有记录仅大小写不同，不能覆盖：{row['id']}"])
            all_orders = {row.id: row for row in session.scalars(select(PurchaseOrder)).all()}
            origins = {row.order_id: row for row in session.scalars(select(OrderOrigin)).all()}
            collisions = (order_ids & set(all_orders)) - set(origins)
            if collisions:
                raise ImportConflict(["单号与模拟/未标记数据冲突，不能覆盖：" + ", ".join(sorted(collisions))])
            existing_shipments = {row.id: row.order_id for row in session.scalars(select(Shipment)).all()}
            for row in data["shipments"]:
                if row["id"] in existing_shipments and existing_shipments[row["id"]] not in order_ids:
                    raise ImportConflict([f"发运编号属于其他订单，不能覆盖：{row['id']}"])
            for table in ("eta_history", "logistics_events", "receipts"):
                existing = {row.id: row.shipment_id for row in session.scalars(select(MODELS[table])).all()}
                for row in data[table]:
                    if row["id"] in existing and existing_shipments[existing[row["id"]]] not in order_ids:
                        raise ImportConflict([f"{table} 编号属于其他订单，不能覆盖：{row['id']}"])
            for row in data["suppliers"]:
                old = session.get(Supplier, row["id"])
                if old and old.name != row["name"]:
                    referenced = any(order.supplier_id == row["id"] and order.id not in order_ids for order in all_orders.values())
                    if referenced:
                        raise ImportConflict([f"供应商 {row['id']} 被其他订单使用，不能修改其名称。"])
            changed = {
                "orders": [(row.id, row.supplier_id, row.material_code, row.qty, str(row.promised_date), str(row.created_at))
                           for row in all_orders.values() if row.id in order_ids],
                "origins": [(key, origins[key].batch_id) for key in sorted(order_ids & set(origins))],
                "suppliers": [(row["id"], session.get(Supplier, row["id"]).name if session.get(Supplier, row["id"]) else None)
                              for row in data["suppliers"]],
            }
            return {"new_orders": len(order_ids - set(all_orders)), "replace_orders": len(order_ids & set(all_orders)),
                    "fingerprint": hashlib.sha256(json.dumps(changed, sort_keys=True, ensure_ascii=False).encode()).hexdigest()}

    def preview(self, files: list[tuple[str, bytes]], as_of: date) -> dict:
        payload, coverage, warnings = parse_uploads(files)
        future_orders = sum(local_datetime(row["created_at"]) >= cutoff_after(as_of) for row in payload["orders"])
        if future_orders:
            warnings.append(f"有 {future_orders} 张订单的创建时间晚于所选快照日，当前分析日期下不可见；请核对日期。")
        existing = self._check_database(payload)
        source_name = ", ".join(Path(name).name for name, _ in files)
        if len(source_name) > 500:
            source_name = source_name[:500]
        token = uuid4().hex
        with self._lock:
            self._previews = {key: value for key, value in self._previews.items() if value.expires_at > time.monotonic()}
            if len(self._previews) >= MAX_PREVIEWS:
                raise ImportValidationError(["待确认预览过多，请完成现有导入或稍后重试。"])
            self._previews[token] = PreparedImport(payload, coverage, as_of, source_name,
                                                  time.monotonic() + PREVIEW_TTL, existing["fingerprint"], warnings)
        return {"preview_id": token, "as_of": as_of.isoformat(), "source_name": source_name,
                "counts": {table: len(rows) for table, rows in payload.items()}, "coverage": coverage,
                "warnings": warnings, "new_orders": existing["new_orders"], "replace_orders": existing["replace_orders"],
                "sample_orders": payload["orders"][:5], "expires_in_seconds": PREVIEW_TTL}

    def commit(self, preview_id: str) -> dict:
        with self._lock:
            prepared = self._previews.get(preview_id)
            if prepared is None or prepared.expires_at < time.monotonic():
                raise ImportConflict(["导入预览已失效，请重新预览。"])
            current = self._check_database(prepared.payload)
            if current["fingerprint"] != prepared.fingerprint:
                raise ImportConflict(["预览后相关数据已变化，请重新预览后再确认。"])
            data = prepared.payload
            ids = [row["id"] for row in data["orders"]]
            batch_id = uuid4().hex
            with Session(self.engine) as session, session.begin():
                shipment_ids = list(session.scalars(select(Shipment.id).where(Shipment.order_id.in_(ids))))
                for model in (Receipt, LogisticsEvent, ETAHistory):
                    if shipment_ids:
                        session.execute(delete(model).where(model.shipment_id.in_(shipment_ids)))
                session.execute(delete(Shipment).where(Shipment.order_id.in_(ids)))
                for table, model in MODELS.items():
                    for row in data[table]:
                        values = {key: date.fromisoformat(value) if key in DATE_FIELDS else
                                  local_datetime(value) if key in DATETIME_FIELDS else value for key, value in row.items()}
                        session.merge(model(**values))
                    session.flush()
                for order_id in ids:
                    session.merge(OrderOrigin(order_id=order_id, batch_id=batch_id, source_name=prepared.source_name,
                                              imported_at=datetime.now(SHANGHAI).replace(tzinfo=None),
                                              as_of=prepared.as_of, coverage=prepared.coverage))
            del self._previews[preview_id]
            return {"status": "imported", "batch_id": batch_id, "order_ids": ids,
                    "as_of": prepared.as_of.isoformat(), "new_orders": current["new_orders"],
                    "replaced_orders": current["replace_orders"], "warnings": prepared.warnings}
